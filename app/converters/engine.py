from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

from PIL import Image
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
from reportlab.pdfgen import canvas

from app.utils.matrix import get_category


def _ffmpeg_convert(source: Path, output: Path, extra_args: list[str] | None = None) -> Path:
    args = ["ffmpeg", "-y", "-i", str(source)]
    if extra_args:
        args.extend(extra_args)
    args.append(str(output))
    subprocess.run(args, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return output


def convert_image(source: Path, output: Path, options: dict) -> Path:
    quality = int(options.get("quality", 85))
    width = options.get("width")
    height = options.get("height")
    with Image.open(source) as img:
        if width and height:
            img = img.resize((int(width), int(height)))
        if output.suffix.lower() in {".jpg", ".jpeg"}:
            img = img.convert("RGB")
        img.save(output, quality=quality)
    return output


def convert_txt_to_pdf(source: Path, output: Path) -> Path:
    c = canvas.Canvas(str(output))
    text = c.beginText(40, 800)
    for line in source.read_text(encoding="utf-8", errors="ignore").splitlines():
        text.textLine(line)
    c.drawText(text)
    c.save()
    return output


def convert_docx(source: Path, output: Path) -> Path:
    doc = Document(source)
    if output.suffix.lower() == ".txt":
        output.write_text("\n".join(p.text for p in doc.paragraphs), encoding="utf-8")
        return output
    if output.suffix.lower() == ".pdf":
        tmp = output.with_suffix(".txt")
        tmp.write_text("\n".join(p.text for p in doc.paragraphs), encoding="utf-8")
        out = convert_txt_to_pdf(tmp, output)
        tmp.unlink(missing_ok=True)
        return out
    doc.save(output)
    return output


def convert_pdf(source: Path, output: Path, options: dict) -> Path:
    reader = PdfReader(str(source))
    writer = PdfWriter()
    page_range = options.get("page_range")
    if page_range:
        selected = []
        for part in str(page_range).split(","):
            if "-" in part:
                a, b = part.split("-", 1)
                selected.extend(range(int(a), int(b) + 1))
            else:
                selected.append(int(part))
        for i in selected:
            if 1 <= i <= len(reader.pages):
                writer.add_page(reader.pages[i - 1])
    else:
        for p in reader.pages:
            writer.add_page(p)
    with output.open("wb") as f:
        writer.write(f)
    return output


def convert_archive(source: Path, output: Path, options: dict) -> Path:
    mode = options.get("archive_mode", "pack")
    if mode == "extract":
        extract_dir = output.with_suffix("")
        extract_dir.mkdir(parents=True, exist_ok=True)
        shutil.unpack_archive(str(source), str(extract_dir))
        packaged = extract_dir.with_suffix(".zip")
        shutil.make_archive(str(packaged.with_suffix("")), "zip", str(extract_dir))
        return packaged

    input_dir = options.get("input_dir")
    if input_dir and Path(input_dir).exists():
        base = Path(input_dir)
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for p in base.rglob("*"):
                if p.is_file():
                    zf.write(p, p.relative_to(base))
        return output

    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.write(source, arcname=source.name)
    return output


def convert_media(source: Path, output: Path, options: dict, category: str) -> Path:
    extra = []
    bitrate = options.get("bitrate")
    frame_rate = options.get("frame_rate")
    sample_rate = options.get("sample_rate")
    if bitrate:
        extra += ["-b:a", str(bitrate)]
    if frame_rate and category == "video":
        extra += ["-r", str(frame_rate)]
    if sample_rate and category == "audio":
        extra += ["-ar", str(sample_rate)]
    return _ffmpeg_convert(source, output, extra)


def convert_file(source: Path, output: Path, target_ext: str, options: dict | str | None = None) -> Path:
    if isinstance(options, str):
        options = json.loads(options) if options.strip() else {}
    options = options or {}
    category = get_category(source.name)

    if category in {"images", "vectors"}:
        return convert_image(source, output, options)
    if category == "documents":
        if source.suffix.lower() == ".txt" and output.suffix.lower() == ".pdf":
            return convert_txt_to_pdf(source, output)
        if source.suffix.lower() == ".docx":
            return convert_docx(source, output)
        if source.suffix.lower() == ".pdf":
            return convert_pdf(source, output, options)
        shutil.copy2(source, output)
        return output
    if category in {"audio", "video"}:
        return convert_media(source, output, options, category)
    if category == "archives":
        return convert_archive(source, output, options)

    shutil.copy2(source, output)
    return output
